"""Iteration 3 regression: user's real /tmp/consumer.csv end-to-end.

- Confirms no runtime/backend regression after native patch-package changes.
- Verifies the CSV pipeline with 10 long multi-line messages (~500 chars each)
  flows through upload -> columns -> validate -> batch -> next -> result
  (with duplicate_ignored) -> counts -> export xlsx.
- Verifies message content is preserved intact (including newlines & full length).
- Verifies /api/sample-template returns xlsx.
- Verifies auth is enforced (403 on missing/invalid token).
"""
import io
import os
import pytest
import requests
import pandas as pd
from openpyxl import load_workbook

# -- Base URL from env / frontend/.env ------------------------------------- #
BASE = os.environ.get("EXPO_PUBLIC_BACKEND_URL")
if not BASE:
    with open("/app/frontend/.env") as f:
        for line in f:
            if line.startswith("EXPO_PUBLIC_BACKEND_URL="):
                BASE = line.split("=", 1)[1].strip()
                break
BASE = BASE.rstrip("/")
API = f"{BASE}/api"
CSV_PATH = "/tmp/consumer.csv"


# -- Fixtures --------------------------------------------------------------- #
@pytest.fixture(scope="module")
def token():
    r = requests.post(f"{API}/gateway/setup", timeout=30)
    assert r.status_code == 200, r.text
    return r.json()["token"]


@pytest.fixture(scope="module")
def h(token):
    return {"X-Gateway-Token": token}


@pytest.fixture(scope="module")
def csv_bytes():
    with open(CSV_PATH, "rb") as f:
        return f.read()


@pytest.fixture(scope="module")
def csv_df(csv_bytes):
    return pd.read_csv(io.BytesIO(csv_bytes), dtype=str, keep_default_na=False)


# -- Auth still enforced ---------------------------------------------------- #
class TestAuthEnforced:
    def test_missing_token(self):
        r = requests.get(f"{API}/dashboard", timeout=30)
        assert r.status_code == 403  # depends on Header(None) -> 403 via our require_token

    def test_invalid_token(self):
        r = requests.get(f"{API}/dashboard", headers={"X-Gateway-Token": "bad"}, timeout=30)
        assert r.status_code == 403

    def test_upload_without_token(self, csv_bytes):
        r = requests.post(f"{API}/files",
                          files={"file": ("consumer.csv", csv_bytes, "text/csv")},
                          timeout=30)
        assert r.status_code == 403

    def test_export_without_token(self):
        # missing token query param -> 403
        r = requests.get(f"{API}/batches/anything/export", timeout=30)
        assert r.status_code == 403


# -- Sample template (public helper) ---------------------------------------- #
class TestSampleTemplate:
    def test_sample_template_xlsx(self):
        r = requests.get(f"{API}/sample-template", timeout=30)
        assert r.status_code == 200
        assert "spreadsheet" in r.headers.get("content-type", "")
        assert "attachment" in r.headers.get("content-disposition", "").lower()
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        header = [c.value for c in ws[1]]
        assert header == ["Mobile Number", "Message"]


# -- Consumer CSV full pipeline -------------------------------------------- #
@pytest.fixture(scope="module")
def uploaded(h, csv_bytes):
    r = requests.post(f"{API}/files", headers=h,
                      files={"file": ("consumer.csv", csv_bytes, "text/csv")},
                      timeout=60)
    assert r.status_code == 200, r.text
    return r.json()


class TestConsumerCsvIngest:
    def test_upload_ok_sheets_csv(self, uploaded):
        assert "file_id" in uploaded
        assert uploaded["sheets"] == ["CSV"]
        assert uploaded["filename"].lower().endswith(".csv")

    def test_columns_and_total_rows(self, h, uploaded, csv_df):
        r = requests.get(f"{API}/files/{uploaded['file_id']}/columns",
                         headers=h, timeout=30)
        assert r.status_code == 200
        c = r.json()
        assert c["sheet"] == "CSV"
        assert c["columns"] == ["Mobile Number", "Message"]
        assert c["total_rows"] == 10
        assert len(c["preview"]) == 10
        # Full multi-line message content preserved intact in first preview row
        first_backend_msg = c["preview"][0]["Message"]
        first_local_msg = csv_df.iloc[0]["Message"]
        assert first_backend_msg == first_local_msg, (
            "Preview did not preserve raw CSV message bytes intact")
        assert len(first_backend_msg) >= 400  # was ~510 chars

    def test_validate_summary_10_valid(self, h, uploaded):
        r = requests.post(f"{API}/files/validate", headers=h, json={
            "file_id": uploaded["file_id"],
            "sheet": "CSV",
            "mobile_column": "Mobile Number",
            "message_column": "Message",
            "country_code": "+91",
        }, timeout=60)
        assert r.status_code == 200, r.text
        v = r.json()
        s = v["summary"]
        assert s == {"total": 10, "valid": 10, "invalid": 0,
                     "duplicate": 0, "empty": 0}, s
        # Numbers normalized to +91
        for row in v["rows"]:
            assert row["mobile"].startswith("+91")
            # 10-digit local -> +91 + 10 digits = 13 char total
            assert len(row["mobile"]) == 13
        # Messages still long -- validator collapses only spaces/tabs; newlines preserved
        assert all(len(row["message"]) > 300 for row in v["rows"])


# -- Full pipeline: batch -> next -> result -> counts -> export ------------ #
@pytest.fixture(scope="module")
def batch(h, uploaded):
    r = requests.post(f"{API}/batches", headers=h, json={
        "file_id": uploaded["file_id"],
        "sheet": "CSV",
        "mobile_column": "Mobile Number",
        "message_column": "Message",
        "country_code": "+91",
    }, timeout=60)
    assert r.status_code == 200, r.text
    b = r.json()
    assert b["total"] == 10
    return b


class TestConsumerCsvPipeline:
    def test_next_atomic_and_result_idempotent(self, h, batch, csv_df):
        # claim one
        r = requests.get(f"{API}/batches/{batch['batch_id']}/next",
                         headers=h, timeout=30)
        assert r.status_code == 200
        payload = r.json()
        job = payload["job"]
        assert job is not None
        assert job["mobile_number"].startswith("+91")
        assert len(job["message"]) > 300  # long message intact in the queue
        msg_id = job["message_id"]

        # complete
        rr = requests.post(f"{API}/jobs/{msg_id}/result", headers=h,
                           json={"status": "sent", "simulated": True}, timeout=30)
        assert rr.status_code == 200
        assert rr.json()["status"] == "SENT"

        # duplicate -> ignored
        rr2 = requests.post(f"{API}/jobs/{msg_id}/result", headers=h,
                            json={"status": "sent", "simulated": True}, timeout=30)
        assert rr2.status_code == 200
        assert rr2.json().get("duplicate_ignored") is True

    def test_process_remaining_9_and_counts(self, h, batch):
        processed = 1  # from previous test
        for _ in range(20):  # safety cap
            r = requests.get(f"{API}/batches/{batch['batch_id']}/next",
                             headers=h, timeout=30)
            assert r.status_code == 200
            job = r.json()["job"]
            if job is None:
                break
            rr = requests.post(f"{API}/jobs/{job['message_id']}/result",
                               headers=h,
                               json={"status": "sent", "simulated": True},
                               timeout=30)
            assert rr.status_code == 200
            processed += 1
        assert processed == 10
        # counts
        r = requests.get(f"{API}/batches/{batch['batch_id']}",
                         headers=h, timeout=30)
        assert r.status_code == 200
        b = r.json()
        assert b["total"] == 10
        assert b["sent"] == 10
        assert b["pending"] == 0
        assert b["sending"] == 0
        assert b["failed"] == 0

    def test_export_xlsx_has_long_messages(self, h, token, batch, csv_df):
        r = requests.get(f"{API}/batches/{batch['batch_id']}/export",
                         params={"token": token}, timeout=60)
        assert r.status_code == 200
        assert "spreadsheet" in r.headers.get("content-type", "")
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        header = [c.value for c in ws[1]]
        assert header == ["Original Row", "Mobile Number", "Message",
                          "Status", "Message ID", "Sent Time", "Error"]
        # Read all rows
        exported = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(exported) == 10
        msg_col = [row[2] for row in exported]
        # All messages should be present with substantial length
        assert all(m and len(m) > 300 for m in msg_col), \
            f"exported message lengths: {[len(m or '') for m in msg_col]}"
        # First exported message should match the (space-collapsed) local first message
        # validator collapses runs of ' '/'\t' only, so equivalent check:
        import re as _re
        expected_first = _re.sub(r"[ \t]+", " ", csv_df.iloc[0]["Message"]).strip()
        # find the row with Original Row = 1
        row1 = next(r for r in exported if r[0] == 1)
        assert row1[2] == expected_first
        assert row1[3] == "SENT"

    def test_export_bad_token(self, batch):
        r = requests.get(f"{API}/batches/{batch['batch_id']}/export",
                         params={"token": "bad"}, timeout=30)
        assert r.status_code == 403
